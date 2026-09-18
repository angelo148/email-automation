from pathlib import Path

import pytest
from openpyxl import Workbook

from app.services.excel_service import (
    ExcelValidationError,
    load_companies_from_excel,
)


def create_test_workbook(
    file_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active

    worksheet.append(
        [
            None,
            "Third Party Integrations for ticket support",
        ]
    )

    worksheet.append(
        [
            None,
            "Module",
            "Third Party",
            "Third Party Group",
            "Email To",
            "Email CC",
        ]
    )

    worksheet.append(
        [
            None,
            "Vouchers",
            "OMTs",
            "OMT",
            "primary@omt.com.lb",
            "primary@omt.com.lb\ncc@omt.com.lb",
        ]
    )

    worksheet.append(
        [
            None,
            None,
            "Alfa",
            "Alfa",
            "first@alfa.com.lb\nsecond@alfa.com.lb",
            "manager@alfa.com.lb\nsupport@alfa.com.lb",
        ]
    )

    worksheet.append(
        [
            None,
            None,
            "Likecard",
            "Likecard",
            "Contact support on their website",
            None,
        ]
    )

    worksheet.append(
        [
            None,
            None,
            "Reloadly",
            "Reloadly",
            None,
            None,
        ]
    )

    workbook.save(file_path)
    workbook.close()


def test_loads_multiline_email_addresses(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "companies.xlsx"

    create_test_workbook(file_path)

    companies = load_companies_from_excel(file_path)

    alfa = next(company for company in companies if company.name == "Alfa")

    assert alfa.module == "Vouchers"
    assert alfa.third_party_group == "Alfa"

    assert [str(address) for address in alfa.to] == [
        "first@alfa.com.lb",
        "second@alfa.com.lb",
    ]

    assert [str(address) for address in alfa.cc] == [
        "manager@alfa.com.lb",
        "support@alfa.com.lb",
    ]


def test_reads_third_party_group(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "companies.xlsx"

    create_test_workbook(file_path)

    companies = load_companies_from_excel(file_path)

    omt = next(company for company in companies if company.name == "OMTs")

    assert omt.third_party_group == "OMT"


def test_inherits_module_from_previous_row(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "companies.xlsx"

    create_test_workbook(file_path)

    companies = load_companies_from_excel(file_path)

    alfa = next(company for company in companies if company.name == "Alfa")

    assert alfa.module == "Vouchers"


def test_duplicate_to_removed_from_cc(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "companies.xlsx"

    create_test_workbook(file_path)

    companies = load_companies_from_excel(file_path)

    omt = next(company for company in companies if company.name == "OMTs")

    assert [str(address) for address in omt.to] == [
        "primary@omt.com.lb",
    ]

    assert [str(address) for address in omt.cc] == [
        "cc@omt.com.lb",
    ]


def test_website_support_is_unavailable(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "companies.xlsx"

    create_test_workbook(file_path)

    companies = load_companies_from_excel(file_path)

    likecard = next(company for company in companies if company.name == "Likecard")

    assert likecard.can_email is False
    assert likecard.to == []

    assert likecard.unavailable_reason == "Contact support on their website"


def test_missing_email_is_unavailable(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "companies.xlsx"

    create_test_workbook(file_path)

    companies = load_companies_from_excel(file_path)

    reloadly = next(company for company in companies if company.name == "Reloadly")

    assert reloadly.can_email is False
    assert reloadly.to == []

    assert reloadly.unavailable_reason == "No email address configured."


def test_missing_third_party_group_raises_error(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "companies.xlsx"

    workbook = Workbook()
    worksheet = workbook.active

    worksheet.append(
        [
            "Module",
            "Third Party",
            "Third Party Group",
            "Email To",
            "Email CC",
        ]
    )

    worksheet.append(
        [
            "Vouchers",
            "Example",
            None,
            "example@example.com",
            None,
        ]
    )

    workbook.save(file_path)
    workbook.close()

    with pytest.raises(
        ExcelValidationError,
        match="Third Party Group is missing",
    ):
        load_companies_from_excel(file_path)


def test_invalid_cc_raises_validation_error(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "companies.xlsx"

    workbook = Workbook()
    worksheet = workbook.active

    worksheet.append(
        [
            "Module",
            "Third Party",
            "Third Party Group",
            "Email To",
            "Email CC",
        ]
    )

    worksheet.append(
        [
            "Vouchers",
            "Example",
            "Example",
            "valid@example.com",
            "not-an-email",
        ]
    )

    workbook.save(file_path)
    workbook.close()

    with pytest.raises(
        ExcelValidationError,
        match="invalid Email CC",
    ):
        load_companies_from_excel(file_path)
