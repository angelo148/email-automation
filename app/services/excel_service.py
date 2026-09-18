import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pydantic import EmailStr, TypeAdapter, ValidationError

from app.schemas.company import CompanyRecord


REQUIRED_COLUMNS = {
    "module",
    "third party",
    "third party group",
    "email to",
    "email cc",
}

HEADER_SEARCH_LIMIT = 50

EMAIL_ADAPTER = TypeAdapter(EmailStr)


class ExcelValidationError(ValueError):
    """Raised when the Excel workbook contains invalid recipient data."""


def _normalize_header(value: Any) -> str:
    """Normalize an Excel header for matching."""

    if value is None:
        return ""

    return str(value).strip().lower()


def _clean_text(value: Any) -> str:
    """Convert an Excel cell into clean single-line text."""

    if value is None:
        return ""

    return " ".join(
        str(value)
        .replace("\r", " ")
        .replace("\n", " ")
        .split()
    ).strip()


def _split_email_values(value: Any) -> list[str]:
    """Split multiple email addresses stored in one Excel cell."""

    if value is None:
        return []

    raw_value = str(value).strip()

    if not raw_value:
        return []

    parts = re.split(r"[\n\r;,]+", raw_value)

    return [
        part.strip()
        for part in parts
        if part.strip()
    ]


def _parse_email_cell(
    value: Any,
) -> tuple[list[EmailStr], list[str]]:
    """Validate email addresses stored in a cell."""

    valid_addresses: list[EmailStr] = []
    invalid_values: list[str] = []

    seen_addresses: set[str] = set()

    for candidate in _split_email_values(value):
        try:
            email = EMAIL_ADAPTER.validate_python(
                candidate
            )

        except ValidationError:
            invalid_values.append(candidate)
            continue

        normalized_email = str(email).lower()

        if normalized_email in seen_addresses:
            continue

        seen_addresses.add(normalized_email)
        valid_addresses.append(email)

    return valid_addresses, invalid_values


def _remove_duplicate_cc_addresses(
    to_addresses: list[EmailStr],
    cc_addresses: list[EmailStr],
) -> list[EmailStr]:
    """Remove CC addresses already present in To."""

    to_lookup = {
        str(address).lower()
        for address in to_addresses
    }

    return [
        address
        for address in cc_addresses
        if str(address).lower() not in to_lookup
    ]


def _find_header_row(
    worksheet: Any,
) -> tuple[int, list[str]]:
    """Find the recipient table header inside the worksheet."""

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            max_row=HEADER_SEARCH_LIMIT,
            values_only=True,
        ),
        start=1,
    ):
        headers = [
            _normalize_header(value)
            for value in row
        ]

        available_headers = {
            header
            for header in headers
            if header
        }

        if REQUIRED_COLUMNS.issubset(
            available_headers
        ):
            return row_number, headers

    required_columns = ", ".join(
        sorted(REQUIRED_COLUMNS)
    )

    raise ExcelValidationError(
        "Could not find the third-party table header. "
        f"Expected columns: {required_columns}"
    )


def _get_unavailable_reason(
    raw_email_to: Any,
) -> str:
    """Return a readable reason for non-email routes."""

    value = _clean_text(raw_email_to)

    if not value:
        return "No email address configured."

    if "website" in value.lower():
        return value

    return (
        "No valid email recipient configured. "
        f"Excel value: {value}"
    )


def load_companies_from_excel(
    file_path: Path,
) -> list[CompanyRecord]:
    """Load third-party recipient configurations from Excel."""

    if not file_path.exists():
        raise FileNotFoundError(
            f"Excel file not found: {file_path}"
        )

    if file_path.suffix.lower() != ".xlsx":
        raise ExcelValidationError(
            "Only .xlsx files are supported."
        )

    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    try:
        worksheet = workbook.active

        header_row_number, headers = _find_header_row(
            worksheet
        )

        column_indexes = {
            header: index
            for index, header in enumerate(headers)
            if header
        }

        companies: list[CompanyRecord] = []
        errors: list[str] = []

        current_module: str | None = None

        rows = worksheet.iter_rows(
            min_row=header_row_number + 1,
            values_only=True,
        )

        for row_number, row in enumerate(
            rows,
            start=header_row_number + 1,
        ):
            module_value = row[
                column_indexes["module"]
            ]

            third_party_value = row[
                column_indexes["third party"]
            ]

            third_party_group_value = row[
                column_indexes["third party group"]
            ]

            email_to_value = row[
                column_indexes["email to"]
            ]

            email_cc_value = row[
                column_indexes["email cc"]
            ]

            module = _clean_text(module_value)

            third_party = _clean_text(
                third_party_value
            )

            third_party_group = _clean_text(
                third_party_group_value
            )

            if module:
                current_module = module

            row_is_empty = not any(
                [
                    module,
                    third_party,
                    third_party_group,
                    _clean_text(email_to_value),
                    _clean_text(email_cc_value),
                ]
            )

            if row_is_empty:
                continue

            if not third_party:
                errors.append(
                    f"Row {row_number}: "
                    "Third Party is missing."
                )
                continue

            if not third_party_group:
                errors.append(
                    f"Row {row_number}: "
                    "Third Party Group is missing."
                )
                continue

            if current_module is None:
                errors.append(
                    f"Row {row_number}: "
                    "Module is missing and cannot "
                    "be inherited from a previous row."
                )
                continue

            (
                to_addresses,
                invalid_to,
            ) = _parse_email_cell(
                email_to_value
            )

            (
                cc_addresses,
                invalid_cc,
            ) = _parse_email_cell(
                email_cc_value
            )

            if to_addresses and invalid_to:
                errors.append(
                    f"Row {row_number}: "
                    "invalid Email To value(s): "
                    f"{', '.join(invalid_to)}"
                )
                continue

            if invalid_cc:
                errors.append(
                    f"Row {row_number}: "
                    "invalid Email CC value(s): "
                    f"{', '.join(invalid_cc)}"
                )
                continue

            cc_addresses = (
                _remove_duplicate_cc_addresses(
                    to_addresses=to_addresses,
                    cc_addresses=cc_addresses,
                )
            )

            can_email = bool(to_addresses)

            unavailable_reason: str | None = None

            if not can_email:
                unavailable_reason = (
                    _get_unavailable_reason(
                        email_to_value
                    )
                )

            try:
                company = CompanyRecord(
                    module=current_module,
                    name=third_party,
                    third_party_group=third_party_group,
                    to=to_addresses,
                    cc=cc_addresses,
                    can_email=can_email,
                    unavailable_reason=(
                        unavailable_reason
                    ),
                    source_row=row_number,
                )

            except ValidationError as exc:
                errors.append(
                    f"Row {row_number}: {exc}"
                )
                continue

            companies.append(company)

        if errors:
            raise ExcelValidationError(
                "Invalid Excel recipient data:\n"
                + "\n".join(errors)
            )

        if not companies:
            raise ExcelValidationError(
                "The Excel workbook contains "
                "no third-party records."
            )

        return companies

    finally:
        workbook.close()